import os
import re
import csv
import io
import hashlib
import secrets
from datetime import date
from contextlib import contextmanager
from pathlib import Path
from typing import Any

import psycopg
from fastapi import Depends, File, Form, Header, HTTPException, UploadFile, status
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from psycopg.rows import dict_row
from pydantic import BaseModel, Field


_raw_database_url = os.getenv("DATABASE_URL", "postgresql://sicv:sicv_dev_password@localhost:5432/sicv")
if _raw_database_url.startswith("postgres://"):
    _raw_database_url = _raw_database_url.replace("postgres://", "postgresql://", 1)
DATABASE_URL = _raw_database_url

ADMIN_API_TOKEN = os.getenv("ADMIN_API_TOKEN", "sicv-admin-local")

_cors_env = os.getenv(
    "CORS_ORIGINS",
    "http://localhost:5173,http://localhost:5174,http://127.0.0.1:5173,http://127.0.0.1:5174",
)
if _cors_env.strip() == "*":
    CORS_ORIGINS = ["*"]
else:
    CORS_ORIGINS = [origin.strip() for origin in _cors_env.split(",") if origin.strip()]

app = FastAPI(
    title="SICV API",
    version="0.1.0",
    description="API MVP para certificacion y verificacion institucional.",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_origin_regex=r"https?://.*\.railway\.app" if CORS_ORIGINS != ["*"] else None,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health")
def health_check():
    return {"status": "ok", "service": "sicv-backend"}


@app.get("/")
def root_check():
    return {"status": "ok", "message": "SICV API activa", "version": "0.1.0"}


DEFAULT_INSTITUTION_ID = "11111111-1111-1111-1111-111111111111"
SYSTEM_ACTOR = "Equipo SICV"
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "/app/uploads"))
QR_DIR = UPLOAD_DIR / "qr"
PUBLIC_VERIFICATION_BASE_URL = os.getenv("PUBLIC_VERIFICATION_BASE_URL", "http://localhost:5173/v").rstrip("/")
ALLOWED_DOCUMENT_MIME = {"application/pdf", "image/jpeg", "image/png", "image/webp"}
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_BULK_ROWS = 500
BULK_TEMPLATE_COLUMNS = [
    "nombres",
    "apellidos",
    "tipo_documento",
    "numero_documento",
    "email",
    "numero_institucional",
]


class HolderCreate(BaseModel):
    tipo_documento: str = Field(min_length=2, max_length=16)
    documento: str = Field(min_length=4, max_length=32)
    nombres: str = Field(min_length=2, max_length=120)
    apellidos: str = Field(min_length=2, max_length=120)


class CertificateCreate(BaseModel):
    holder_id: str
    type_id: str
    numero_institucional: str = Field(min_length=4, max_length=64)
    titulo: str = Field(min_length=4, max_length=200)
    programa: str | None = Field(default=None, max_length=160)
    fecha_emision: date
    vence_en: date | None = None
    origen: str = "NUEVO"


class StateAction(BaseModel):
    actor: str = Field(default=SYSTEM_ACTOR, min_length=2, max_length=120)
    reason: str | None = Field(default=None, max_length=300)


class PolicyCreate(BaseModel):
    descripcion: str = Field(min_length=8, max_length=300)
    mostrar_codigo: bool = True
    mostrar_nombre: bool = True
    mostrar_documento: bool = False
    mostrar_fechas: bool = True
    mostrar_programa: bool = True
    permitir_documento_completo: bool = False
    actor: str = Field(default=SYSTEM_ACTOR, min_length=2, max_length=120)


class InstitutionSettingsUpdate(BaseModel):
    razon_social: str = Field(min_length=2, max_length=160)
    dominio_verificacion: str = Field(min_length=3, max_length=255)


@app.on_event("startup")
def ensure_runtime_schema() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    QR_DIR.mkdir(parents=True, exist_ok=True)
    with db_cursor() as cur:
        cur.execute("create extension if not exists pgcrypto;")

        cur.execute(
            """
            create table if not exists institucion (
              id uuid primary key default gen_random_uuid(),
              razon_social text not null,
              dominio_verificacion text not null,
              activo boolean not null default true,
              created_at timestamptz not null default now()
            )
            """
        )
        cur.execute(
            """
            insert into institucion (id, razon_social, dominio_verificacion)
            values (%s, 'Institucion Educativa Demo', 'localhost')
            on conflict (id) do nothing
            """,
            (DEFAULT_INSTITUTION_ID,),
        )

        cur.execute(
            """
            create table if not exists titular (
              id uuid primary key default gen_random_uuid(),
              institucion_id uuid not null references institucion(id) on delete restrict,
              tipo_documento varchar(16) not null,
              documento_normalizado varchar(64) not null,
              documento_mascara varchar(32) not null,
              nombres text not null,
              apellidos text not null,
              created_at timestamptz not null default now(),
              unique (institucion_id, tipo_documento, documento_normalizado)
            )
            """
        )

        cur.execute(
            """
            create table if not exists tipo_certificado (
              id uuid primary key default gen_random_uuid(),
              institucion_id uuid not null references institucion(id) on delete restrict,
              codigo varchar(32) not null,
              nombre text not null,
              requiere_vigencia boolean not null default false,
              activo boolean not null default true,
              unique (institucion_id, codigo)
            )
            """
        )
        cur.execute(
            """
            insert into tipo_certificado (id, institucion_id, codigo, nombre, requiere_vigencia)
            values
              ('22222222-2222-2222-2222-222222222221', '11111111-1111-1111-1111-111111111111', 'CAP', 'Certificado de capacitacion', false),
              ('22222222-2222-2222-2222-222222222222', '11111111-1111-1111-1111-111111111111', 'HIST', 'Constancia historica de estudios', false),
              ('22222222-2222-2222-2222-222222222223', '11111111-1111-1111-1111-111111111111', 'TEC', 'Certificado tecnico', true)
            on conflict (institucion_id, codigo) do nothing;
            """
        )

        cur.execute(
            """
            create table if not exists certificado (
              id uuid primary key default gen_random_uuid(),
              institucion_id uuid not null references institucion(id) on delete restrict,
              titular_id uuid not null references titular(id) on delete restrict,
              tipo_id uuid not null references tipo_certificado(id) on delete restrict,
              numero_institucional varchar(64) not null,
              token_publico varchar(128) unique,
              estado varchar(24) not null,
              origen varchar(16) not null,
              titulo text not null,
              programa text,
              fecha_emision date not null,
              vence_en date,
              revisor text,
              created_at timestamptz not null default now(),
              updated_at timestamptz not null default now(),
              unique (institucion_id, numero_institucional),
              check (estado in ('BORRADOR', 'EN_REVISION', 'OBSERVADO', 'APROBADO', 'PUBLICADO', 'REVOCADO', 'VENCIDO', 'REEMPLAZADO')),
              check (origen in ('HISTORICO', 'NUEVO')),
              check (vence_en is null or vence_en >= fecha_emision)
            )
            """
        )
        cur.execute("create index if not exists ix_certificado_token_publico on certificado(token_publico);")
        cur.execute("create index if not exists ix_certificado_estado_created on certificado(estado, created_at desc);")

        cur.execute(
            """
            create table if not exists politica_divulgacion (
              id uuid primary key default gen_random_uuid(),
              institucion_id uuid not null references institucion(id) on delete restrict,
              tipo_id uuid references tipo_certificado(id) on delete restrict,
              version integer not null,
              descripcion text not null,
              activa boolean not null default false,
              mostrar_codigo boolean not null default true,
              mostrar_nombre boolean not null default false,
              mostrar_documento boolean not null default false,
              mostrar_fechas boolean not null default true,
              mostrar_programa boolean not null default true,
              permitir_documento_completo boolean not null default false,
              approved_at timestamptz,
              created_at timestamptz not null default now()
            )
            """
        )
        cur.execute(
            """
            create unique index if not exists uq_politica_activa_tipo
              on politica_divulgacion (institucion_id, tipo_id)
              where activa = true and tipo_id is not null;
            """
        )
        cur.execute(
            """
            create unique index if not exists uq_politica_activa_general
              on politica_divulgacion (institucion_id)
              where activa = true and tipo_id is null;
            """
        )
        cur.execute(
            """
            insert into politica_divulgacion (
              id, institucion_id, tipo_id, version, descripcion, activa,
              mostrar_codigo, mostrar_nombre, mostrar_documento, mostrar_fechas,
              mostrar_programa, permitir_documento_completo, approved_at
            )
            values
              (
                '44444444-4444-4444-4444-444444444441',
                '11111111-1111-1111-1111-111111111111',
                null,
                1,
                'Codigo, titulo, institucion, fecha, programa y estado. Documento completo oculto por defecto.',
                true, true, true, false, true, true, false, now()
              )
            on conflict do nothing;
            """
        )

        cur.execute(
            """
            create table if not exists documento_version (
              id uuid primary key default gen_random_uuid(),
              certificado_id uuid not null references certificado(id) on delete restrict,
              version integer not null,
              object_key text not null,
              sha256 char(64) not null,
              mime varchar(128) not null,
              size_bytes bigint not null,
              es_definitivo boolean not null default false,
              created_at timestamptz not null default now(),
              unique (certificado_id, version)
            )
            """
        )
        cur.execute(
            """
            create unique index if not exists uq_documento_definitivo
              on documento_version (certificado_id)
              where es_definitivo = true;
            """
        )

        cur.execute(
            """
            create table if not exists certificado_qr (
              certificado_id uuid primary key references certificado(id) on delete cascade,
              verification_url text not null,
              object_key text not null,
              sha256 char(64) not null,
              mime varchar(64) not null default 'image/svg+xml',
              created_at timestamptz not null default now()
            )
            """
        )

        cur.execute(
            """
            create table if not exists carga_masiva (
              id uuid primary key default gen_random_uuid(),
              institucion_id uuid not null references institucion(id) on delete restrict,
              tipo_id uuid references tipo_certificado(id) on delete restrict,
              titulo text not null,
              programa text,
              fecha_emision date not null,
              vence_en date,
              origen varchar(16) not null,
              total_rows integer not null default 0,
              valid_rows integer not null default 0,
              error_rows integer not null default 0,
              created_count integer not null default 0,
              skipped_count integer not null default 0,
              actor text not null,
              status varchar(24) not null default 'COMPLETADO',
              created_at timestamptz not null default now()
            )
            """
        )

        cur.execute(
            """
            create table if not exists carga_masiva_detalle (
              id uuid primary key default gen_random_uuid(),
              carga_id uuid not null references carga_masiva(id) on delete cascade,
              row_number integer not null,
              certificado_id uuid references certificado(id) on delete set null,
              numero_institucional varchar(64),
              titular text,
              documento text,
              status varchar(24) not null,
              errors text,
              warnings text,
              verification_url text,
              created_at timestamptz not null default now()
            )
            """
        )

        cur.execute(
            """
            create table if not exists cola_revision (
              id uuid primary key default gen_random_uuid(),
              id_publico varchar(32) not null unique,
              titular text not null,
              titulo text not null,
              origen varchar(16) not null,
              estado varchar(24) not null,
              actualizado text not null,
              riesgo text not null,
              orden integer not null default 0
            )
            """
        )

        cur.execute(
            """
            create table if not exists evento_auditoria (
              id uuid primary key default gen_random_uuid(),
              accion varchar(64) not null,
              objetivo text not null,
              actor text not null,
              creado_en timestamptz not null default now()
            )
            """
        )


@contextmanager
def db_cursor():
    with psycopg.connect(DATABASE_URL, row_factory=dict_row) as conn:
        with conn.cursor() as cur:
            yield cur


def require_admin(x_admin_token: str | None = Header(default=None)) -> None:
    if not x_admin_token or not secrets.compare_digest(x_admin_token, ADMIN_API_TOKEN):
        raise HTTPException(status_code=403, detail="admin_required")


def fetch_all(query: str, params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    with db_cursor() as cur:
        cur.execute(query, params)
        return list(cur.fetchall())


def fetch_one(query: str, params: tuple[Any, ...] = ()) -> dict[str, Any] | None:
    with db_cursor() as cur:
        cur.execute(query, params)
        return cur.fetchone()


def execute_one(query: str, params: tuple[Any, ...] = ()) -> dict[str, Any] | None:
    with db_cursor() as cur:
        cur.execute(query, params)
        return cur.fetchone()


def normalize_document(value: str) -> str:
    return re.sub(r"\W+", "", value).upper()


def mask_document(value: str) -> str:
    clean = normalize_document(value)
    return f"****{clean[-4:]}" if len(clean) >= 4 else "****"


def generate_public_token() -> str:
    return secrets.token_urlsafe(18)


def get_verification_base_url(cur: Any | None = None) -> str:
    raw = None
    try:
        if cur is not None:
            cur.execute(
                "select dominio_verificacion from institucion where id = %s",
                (DEFAULT_INSTITUTION_ID,),
            )
            row = cur.fetchone()
        else:
            row = fetch_one(
                "select dominio_verificacion from institucion where id = %s",
                (DEFAULT_INSTITUTION_ID,),
            )
        if row and row.get("dominio_verificacion"):
            raw = str(row["dominio_verificacion"]).strip()
    except Exception:
        raw = None

    if not raw or raw in ("localhost", "127.0.0.1"):
        if PUBLIC_VERIFICATION_BASE_URL and ("localhost" in PUBLIC_VERIFICATION_BASE_URL or "127.0.0.1" in PUBLIC_VERIFICATION_BASE_URL):
            return PUBLIC_VERIFICATION_BASE_URL

    if not raw:
        return PUBLIC_VERIFICATION_BASE_URL

    raw = raw.rstrip("/")
    if raw.startswith("http://") or raw.startswith("https://"):
        return raw if raw.endswith("/v") else f"{raw}/v"

    scheme = "http" if ("localhost" in raw or "127.0.0.1" in raw) else "https"
    return f"{scheme}://{raw}/v"


def verification_url_for_token(token: str, cur: Any | None = None) -> str:
    base_url = get_verification_base_url(cur)
    return f"{base_url}/{token}"


def generate_qr_svg(certificado_id: str, token: str) -> dict[str, Any]:
    import qrcode
    import qrcode.image.svg

    QR_DIR.mkdir(parents=True, exist_ok=True)
    verification_url = verification_url_for_token(token)
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(verification_url)
    qr.make(fit=True)
    image = qr.make_image(image_factory=qrcode.image.svg.SvgPathImage)
    svg_bytes = image.to_string(encoding="unicode").encode("utf-8")
    object_key = f"qr/{certificado_id}.svg"
    destination = UPLOAD_DIR / object_key
    destination.write_bytes(svg_bytes)
    digest = hashlib.sha256(svg_bytes).hexdigest()

    return upsert_qr_metadata(certificado_id, verification_url, object_key, digest)


def upsert_qr_metadata(certificado_id: str, verification_url: str, object_key: str, digest: str) -> dict[str, Any]:
    with db_cursor() as cur:
        cur.execute(
            """
            insert into certificado_qr (certificado_id, verification_url, object_key, sha256, mime)
            values (%s, %s, %s, %s, 'image/svg+xml')
            on conflict (certificado_id) do update
            set verification_url = excluded.verification_url,
                object_key = excluded.object_key,
                sha256 = excluded.sha256,
                mime = excluded.mime,
                created_at = now()
            returning verification_url, object_key, sha256, mime
            """,
            (certificado_id, verification_url, object_key, digest),
        )
        return cur.fetchone()


def generate_qr_svg_in_cursor(cur: Any, certificado_id: str, token: str) -> dict[str, Any]:
    import qrcode
    import qrcode.image.svg

    QR_DIR.mkdir(parents=True, exist_ok=True)
    verification_url = verification_url_for_token(token, cur=cur)
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(verification_url)
    qr.make(fit=True)
    image = qr.make_image(image_factory=qrcode.image.svg.SvgPathImage)
    svg_bytes = image.to_string(encoding="unicode").encode("utf-8")
    object_key = f"qr/{certificado_id}.svg"
    destination = UPLOAD_DIR / object_key
    destination.write_bytes(svg_bytes)
    digest = hashlib.sha256(svg_bytes).hexdigest()
    cur.execute(
        """
        insert into certificado_qr (certificado_id, verification_url, object_key, sha256, mime)
        values (%s, %s, %s, %s, 'image/svg+xml')
        on conflict (certificado_id) do update
        set verification_url = excluded.verification_url,
            object_key = excluded.object_key,
            sha256 = excluded.sha256,
            mime = excluded.mime,
            created_at = now()
        returning verification_url, object_key, sha256, mime
        """,
        (certificado_id, verification_url, object_key, digest),
    )
    return cur.fetchone()


def ensure_qr_for_certificate(certificate: dict[str, Any]) -> dict[str, Any]:
    token = certificate["token_publico"] or generate_public_token()
    if not certificate["token_publico"]:
        execute_one(
            "update certificado set token_publico = %s where id = %s returning id",
            (token, certificate["id"]),
        )
        certificate["token_publico"] = token

    existing = fetch_one(
        "select verification_url, object_key, sha256, mime from certificado_qr where certificado_id = %s",
        (certificate["id"],),
    )
    if existing is not None:
        return existing
    return generate_qr_svg(str(certificate["id"]), token)


def audit(action: str, target: str, actor: str = SYSTEM_ACTOR) -> None:
    execute_one(
        """
        insert into evento_auditoria (accion, objetivo, actor)
        values (%s, %s, %s)
        returning id
        """,
        (action, target, actor),
    )


def audit_in_cursor(cur: Any, action: str, target: str, actor: str = SYSTEM_ACTOR) -> None:
    cur.execute(
        """
        insert into evento_auditoria (accion, objetivo, actor)
        values (%s, %s, %s)
        returning id
        """,
        (action, target, actor),
    )


def clean_uploads() -> None:
    if not UPLOAD_DIR.exists():
        return
    for path in UPLOAD_DIR.rglob("*"):
        if path.is_file():
            path.unlink(missing_ok=True)


def certificate_projection(
    where_clause: str = "",
    params: tuple[Any, ...] = (),
    limit: int | None = None,
) -> list[dict[str, Any]]:
    limit_clause = f"limit {limit}" if limit else ""
    return fetch_all(
        f"""
        select
          c.id,
          c.token_publico as token,
          c.numero_institucional as code,
          concat(t.nombres, ' ', t.apellidos) as holder,
          t.id as holder_id,
          concat(t.tipo_documento, ' ', t.documento_mascara) as masked_document,
          c.titulo as title,
          c.programa as program,
          i.razon_social as issuer,
          c.fecha_emision as issued_on,
          c.vence_en as expires_on,
          c.estado as status,
          c.origen as origin,
          c.revisor as reviewer,
          coalesce(pd.descripcion, 'Politica institucional vigente') as disclosure,
          exists (
            select 1 from documento_version dv
            where dv.certificado_id = c.id and dv.es_definitivo = true
          ) as has_document,
          (cq.certificado_id is not null) as has_qr,
          cq.verification_url as qr_url,
          c.created_at
        from certificado c
        join institucion i on i.id = c.institucion_id
        join titular t on t.id = c.titular_id
        join tipo_certificado tc on tc.id = c.tipo_id
        left join politica_divulgacion pd on pd.institucion_id = c.institucion_id
          and (pd.tipo_id = c.tipo_id or pd.tipo_id is null)
          and pd.activa = true
        left join certificado_qr cq on cq.certificado_id = c.id
        {where_clause}
        order by c.created_at desc
        {limit_clause}
        """,
        params,
    )


def get_certificate_or_404(certificate_id: str) -> dict[str, Any]:
    certificate = fetch_one(
        "select * from certificado where id = %s",
        (certificate_id,),
    )
    if certificate is None:
        raise HTTPException(status_code=404, detail="certificate_not_found")
    return certificate


def transition_certificate(
    certificate_id: str,
    allowed_states: set[str],
    next_state: str,
    action: str,
    actor: str,
    reason: str | None = None,
) -> dict[str, Any]:
    certificate = get_certificate_or_404(certificate_id)
    if certificate["estado"] not in allowed_states:
        raise HTTPException(status_code=409, detail="invalid_state_transition")

    target = certificate["numero_institucional"]
    updated = execute_one(
        """
        update certificado
        set estado = %s, revisor = %s, updated_at = now()
        where id = %s
        returning id
        """,
        (next_state, actor, certificate_id),
    )
    if updated is None:
        raise HTTPException(status_code=404, detail="certificate_not_found")

    audit(action, f"{target}: {reason}" if reason else target, actor)
    return certificate_projection("where c.id = %s", (certificate_id,))[0]


def normalize_bulk_headers(headers: list[str]) -> dict[str, str]:
    return {header.strip().lower(): header for header in headers if header}


def validate_bulk_headers(headers: dict[str, str]) -> None:
    missing = [column for column in BULK_TEMPLATE_COLUMNS[:4] if column not in headers]
    if missing:
        raise HTTPException(status_code=422, detail=f"missing_columns:{','.join(missing)}")


def read_bulk_csv(raw: bytes) -> list[dict[str, Any]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=422, detail="csv_must_be_utf8") from exc

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(status_code=422, detail="csv_header_required")

    normalized_headers = normalize_bulk_headers(reader.fieldnames)
    validate_bulk_headers(normalized_headers)
    rows = []
    for index, row in enumerate(reader, start=2):
        if len(rows) >= MAX_BULK_ROWS:
            raise HTTPException(status_code=413, detail="bulk_row_limit_exceeded")
        normalized = {
            column: (row.get(original, "") or "").strip()
            for column, original in normalized_headers.items()
        }
        if not any(normalized.values()):
            continue
        rows.append({"row": index, "data": normalized})
    return rows


def read_bulk_xlsx(raw: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="xlsx_invalid_or_corrupt") from exc

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(status_code=422, detail="xlsx_header_required") from exc

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    normalized_headers = normalize_bulk_headers(headers)
    validate_bulk_headers(normalized_headers)
    rows = []
    for index, values in enumerate(rows_iter, start=2):
        if len(rows) >= MAX_BULK_ROWS:
            raise HTTPException(status_code=413, detail="bulk_row_limit_exceeded")
        raw_row = {
            headers[column_index]: "" if value is None else str(value).strip()
            for column_index, value in enumerate(values)
            if column_index < len(headers) and headers[column_index]
        }
        normalized = {
            column: (raw_row.get(original, "") or "").strip()
            for column, original in normalized_headers.items()
        }
        if not any(normalized.values()):
            continue
        rows.append({"row": index, "data": normalized})
    return rows


def read_bulk_file(file: UploadFile) -> list[dict[str, Any]]:
    raw = file.file.read()
    if not raw:
        raise HTTPException(status_code=422, detail="empty_file")
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="bulk_file_too_large")

    filename = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()
    if filename.endswith(".xlsx") or content_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    }:
        rows = read_bulk_xlsx(raw)
    elif filename.endswith(".csv") or content_type in {"text/csv", "application/vnd.ms-excel", "text/plain"}:
        rows = read_bulk_csv(raw)
    else:
        raise HTTPException(status_code=415, detail="unsupported_bulk_file_type")

    if not rows:
        raise HTTPException(status_code=422, detail="bulk_file_without_data")
    return rows


def automatic_certificate_suffix(row_number: int, data: dict[str, str], prefix: str) -> str:
    base = "|".join(
        [
            DEFAULT_INSTITUTION_ID,
            prefix,
            str(row_number),
            data.get("tipo_documento", ""),
            normalize_document(data.get("numero_documento", "")),
            data.get("nombres", ""),
            data.get("apellidos", ""),
        ]
    )
    digest = hashlib.sha256(base.encode("utf-8")).hexdigest().upper()
    return digest[:8]


def build_certificate_number(
    row_number: int,
    data: dict[str, str],
    prefix: str,
    start: int,
    mode: str,
) -> str:
    explicit = data.get("numero_institucional", "").strip()
    if explicit:
        return explicit
    if mode == "CORRELATIVO":
        return f"{prefix}{start + row_number - 1:06d}"
    return f"{prefix}{automatic_certificate_suffix(row_number, data, prefix)}"


def validate_bulk_rows(
    rows: list[dict[str, Any]],
    type_id: str,
    titulo: str,
    fecha_emision: date,
    vence_en: date | None,
    origen: str,
    number_prefix: str,
    number_start: int,
    number_mode: str,
) -> dict[str, Any]:
    if vence_en and vence_en < fecha_emision:
        raise HTTPException(status_code=422, detail="expires_before_issue_date")
    if origen not in {"NUEVO", "HISTORICO"}:
        raise HTTPException(status_code=422, detail="invalid_origin")
    if number_start < 1:
        raise HTTPException(status_code=422, detail="invalid_number_start")
    if number_mode not in {"AUTOMATICO", "CORRELATIVO"}:
        raise HTTPException(status_code=422, detail="invalid_number_mode")

    certificate_type = fetch_one(
        """
        select id, nombre
        from tipo_certificado
        where id = %s and institucion_id = %s and activo = true
        """,
        (type_id, DEFAULT_INSTITUTION_ID),
    )
    if certificate_type is None:
        raise HTTPException(status_code=422, detail="invalid_certificate_type")

    seen_documents: dict[tuple[str, str], int] = {}
    seen_numbers: dict[str, int] = {}
    certificate_numbers: list[str] = []
    document_keys: list[tuple[str, str]] = []
    preview_rows = []

    for position, item in enumerate(rows, start=1):
        row_number = item["row"]
        data = item["data"]
        errors = []
        warnings = []
        nombres = data.get("nombres", "").strip()
        apellidos = data.get("apellidos", "").strip()
        tipo_documento = data.get("tipo_documento", "DNI").strip().upper()
        numero_documento = data.get("numero_documento", "").strip()
        normalized_document = normalize_document(numero_documento)
        certificate_number = build_certificate_number(position, data, number_prefix, number_start, number_mode).strip()

        if len(nombres) < 2:
            errors.append("nombres_requeridos")
        if len(apellidos) < 2:
            errors.append("apellidos_requeridos")
        if len(tipo_documento) < 2 or len(tipo_documento) > 16:
            errors.append("tipo_documento_invalido")
        if len(normalized_document) < 4:
            errors.append("numero_documento_invalido")
        if len(certificate_number) < 4 or len(certificate_number) > 64:
            errors.append("numero_institucional_invalido")

        document_key = (tipo_documento, normalized_document)
        if normalized_document:
            if document_key in seen_documents:
                warnings.append(f"documento_repetido_en_fila_{seen_documents[document_key]}")
            seen_documents[document_key] = row_number
            document_keys.append(document_key)

        if certificate_number:
            normalized_number = certificate_number.upper()
            if normalized_number in seen_numbers:
                errors.append(f"numero_institucional_repetido_en_fila_{seen_numbers[normalized_number]}")
            seen_numbers[normalized_number] = row_number
            certificate_numbers.append(certificate_number)

        preview_rows.append(
            {
                "row": row_number,
                "nombres": nombres,
                "apellidos": apellidos,
                "tipo_documento": tipo_documento,
                "numero_documento": numero_documento,
                "documento_mascara": mask_document(numero_documento) if numero_documento else "",
                "numero_institucional": certificate_number,
                "errors": errors,
                "warnings": warnings,
                "status": "VALIDO" if not errors else "ERROR",
            }
        )

    existing_numbers = set()
    if certificate_numbers:
        existing = fetch_all(
            """
            select upper(numero_institucional) as numero
            from certificado
            where institucion_id = %s and upper(numero_institucional) = any(%s)
            """,
            (DEFAULT_INSTITUTION_ID, [number.upper() for number in certificate_numbers]),
        )
        existing_numbers = {row["numero"] for row in existing}

    existing_documents = set()
    for tipo_documento, normalized_document in document_keys:
        holder = fetch_one(
            """
            select id
            from titular
            where institucion_id = %s and tipo_documento = %s and documento_normalizado = %s
            """,
            (DEFAULT_INSTITUTION_ID, tipo_documento, normalized_document),
        )
        if holder:
            existing_documents.add((tipo_documento, normalized_document))

    for row in preview_rows:
        if row["numero_institucional"].upper() in existing_numbers:
            row["errors"].append("numero_institucional_ya_existe")
        document_key = (row["tipo_documento"], normalize_document(row["numero_documento"]))
        if document_key in existing_documents:
            row["warnings"].append("titular_existente_se_reutilizara")
        row["status"] = "VALIDO" if not row["errors"] else "ERROR"

    valid_rows = sum(1 for row in preview_rows if not row["errors"])
    error_rows = len(preview_rows) - valid_rows
    return {
        "certificate_type": certificate_type,
        "total_rows": len(preview_rows),
        "valid_rows": valid_rows,
        "error_rows": error_rows,
        "rows": preview_rows,
    }


def import_bulk_rows(
    preview: dict[str, Any],
    type_id: str,
    titulo: str,
    programa: str | None,
    fecha_emision: date,
    vence_en: date | None,
    origen: str,
    actor: str,
    allow_skip_invalid: bool,
) -> dict[str, Any]:
    if preview["error_rows"] and not allow_skip_invalid:
        raise HTTPException(status_code=409, detail="bulk_preview_has_errors")

    imported = []
    skipped = []
    with db_cursor() as cur:
        cur.execute(
            """
            insert into carga_masiva (
              institucion_id, tipo_id, titulo, programa, fecha_emision, vence_en, origen,
              total_rows, valid_rows, error_rows, actor, status
            )
            values (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'COMPLETADO')
            returning id
            """,
            (
                DEFAULT_INSTITUTION_ID,
                type_id,
                titulo.strip(),
                programa.strip() if programa else None,
                fecha_emision,
                vence_en,
                origen,
                preview["total_rows"],
                preview["valid_rows"],
                preview["error_rows"],
                actor,
            ),
        )
        bulk = cur.fetchone()

        for row in preview["rows"]:
            if row["errors"]:
                skipped.append(row)
                cur.execute(
                    """
                    insert into carga_masiva_detalle (
                      carga_id, row_number, numero_institucional, titular, documento,
                      status, errors, warnings
                    )
                    values (%s, %s, %s, %s, %s, 'OMITIDO', %s, %s)
                    """,
                    (
                        bulk["id"],
                        row["row"],
                        row["numero_institucional"],
                        f"{row['nombres']} {row['apellidos']}".strip(),
                        f"{row['tipo_documento']} {row['documento_mascara']}".strip(),
                        ", ".join(row["errors"]),
                        ", ".join(row["warnings"]),
                    ),
                )
                continue

            normalized_document = normalize_document(row["numero_documento"])
            cur.execute(
                """
                select id
                from titular
                where institucion_id = %s and tipo_documento = %s and documento_normalizado = %s
                """,
                (DEFAULT_INSTITUTION_ID, row["tipo_documento"], normalized_document),
            )
            holder = cur.fetchone()
            if holder is None:
                cur.execute(
                    """
                    insert into titular (
                      institucion_id, tipo_documento, documento_normalizado,
                      documento_mascara, nombres, apellidos
                    )
                    values (%s, %s, %s, %s, %s, %s)
                    returning id
                    """,
                    (
                        DEFAULT_INSTITUTION_ID,
                        row["tipo_documento"],
                        normalized_document,
                        mask_document(row["numero_documento"]),
                        row["nombres"],
                        row["apellidos"],
                    ),
                )
                holder = cur.fetchone()

            token = generate_public_token()
            cur.execute(
                """
                insert into certificado (
                  institucion_id, titular_id, tipo_id, numero_institucional,
                  token_publico, estado, origen, titulo, programa, fecha_emision, vence_en
                )
                values (%s, %s, %s, %s, %s, 'BORRADOR', %s, %s, %s, %s, %s)
                returning id, numero_institucional, token_publico
                """,
                (
                    DEFAULT_INSTITUTION_ID,
                    holder["id"],
                    type_id,
                    row["numero_institucional"],
                    token,
                    origen,
                    titulo.strip(),
                    programa.strip() if programa else None,
                    fecha_emision,
                    vence_en,
                ),
            )
            certificate = cur.fetchone()
            qr = generate_qr_svg_in_cursor(cur, str(certificate["id"]), token)
            cur.execute(
                """
                insert into carga_masiva_detalle (
                  carga_id, row_number, certificado_id, numero_institucional, titular,
                  documento, status, errors, warnings, verification_url
                )
                values (%s, %s, %s, %s, %s, %s, 'CREADO', null, %s, %s)
                """,
                (
                    bulk["id"],
                    row["row"],
                    certificate["id"],
                    certificate["numero_institucional"],
                    f"{row['nombres']} {row['apellidos']}".strip(),
                    f"{row['tipo_documento']} {row['documento_mascara']}".strip(),
                    ", ".join(row["warnings"]),
                    qr["verification_url"],
                ),
            )
            imported.append(
                {
                    "row": row["row"],
                    "certificate_id": certificate["id"],
                    "numero_institucional": certificate["numero_institucional"],
                    "token": certificate["token_publico"],
                    "verification_url": qr["verification_url"],
                }
            )

        cur.execute(
            """
            update carga_masiva
            set created_count = %s, skipped_count = %s
            where id = %s
            """,
            (len(imported), len(skipped), bulk["id"]),
        )
        audit_in_cursor(cur, "CARGA_MASIVA", f"{len(imported)} certificados creados", actor)
        if imported:
            audit_in_cursor(cur, "GENERAR_QR_MASIVO", f"{len(imported)} QR generados", actor)

    return {
        "bulk_id": bulk["id"],
        "created": len(imported),
        "skipped": len(skipped),
        "results": imported,
        "skipped_rows": skipped,
    }


def bulk_action_result(bulk_id: str, actor: str) -> dict[str, Any]:
    with db_cursor() as cur:
        cur.execute(
            """
            update certificado c
            set estado = 'EN_REVISION', revisor = %s, updated_at = now()
            from carga_masiva_detalle d
            where d.carga_id = %s
              and d.certificado_id = c.id
              and c.estado in ('BORRADOR', 'OBSERVADO')
            returning c.id, c.numero_institucional
            """,
            (actor, bulk_id),
        )
        updated = list(cur.fetchall())
        if updated:
            audit_in_cursor(cur, "ENVIAR_REVISION_MASIVO", f"{len(updated)} certificados", actor)
        return {
            "updated": len(updated),
            "certificates": updated,
        }


@app.get("/health/live")
def live():
    return {"status": "ok"}


@app.get("/health/ready")
def ready():
    try:
        row = fetch_one("select 1 as ok")
    except Exception as exc:  # pragma: no cover - defensive runtime endpoint
        raise HTTPException(status_code=503, detail="database_unavailable") from exc
    return {"status": "ready", "database": row["ok"] == 1}


@app.get("/api/v1/public/status")
def public_status():
    return {"status": "available", "service": "SICV"}


@app.get("/api/v1/public/certificates/{token}")
def verify_certificate(token: str):
    certificate = fetch_one(
        """
        select
          c.token_publico as token,
          c.id as id,
          c.numero_institucional as code,
          c.estado as status,
          c.titulo as title,
          c.programa as program,
          c.origen as origin,
          c.fecha_emision as issued_on,
          c.vence_en as expires_on,
          i.razon_social as issuer,
          i.dominio_verificacion as issuer_domain,
          case when pd.mostrar_nombre then concat(t.nombres, ' ', t.apellidos) else null end as holder,
          case when pd.mostrar_documento then concat(t.tipo_documento, ' ', t.documento_mascara) else 'No publico' end as masked_document,
          pd.descripcion as disclosure,
          pd.permitir_documento_completo as document_public,
          exists (
            select 1 from documento_version dv
            where dv.certificado_id = c.id and dv.es_definitivo = true
          ) as has_document
        from certificado c
        join institucion i on i.id = c.institucion_id
        join titular t on t.id = c.titular_id
        join tipo_certificado tc on tc.id = c.tipo_id
        join politica_divulgacion pd on pd.institucion_id = c.institucion_id
          and (pd.tipo_id = c.tipo_id or pd.tipo_id is null)
          and pd.activa = true
        where (c.token_publico = %s or c.numero_institucional = %s)
          and c.estado in ('PUBLICADO', 'REVOCADO', 'VENCIDO', 'REEMPLAZADO')
        order by pd.tipo_id nulls last
        limit 1
        """,
        (token, token),
    )
    if certificate is None:
        raise HTTPException(status_code=404, detail="certificate_not_found")
    return certificate


@app.get("/api/v1/public/certificates/{token}/document")
def public_certificate_document(token: str):
    document = fetch_one(
        """
        select dv.object_key, dv.mime
        from certificado c
        join politica_divulgacion pd on pd.institucion_id = c.institucion_id
          and (pd.tipo_id = c.tipo_id or pd.tipo_id is null)
          and pd.activa = true
        join documento_version dv on dv.certificado_id = c.id and dv.es_definitivo = true
        where c.token_publico = %s
          and c.estado = 'PUBLICADO'
          and pd.permitir_documento_completo = true
        order by pd.tipo_id nulls last
        limit 1
        """,
        (token,),
    )
    if document is None:
        raise HTTPException(status_code=404, detail="document_not_available")
    path = UPLOAD_DIR / document["object_key"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="document_file_missing")
    return FileResponse(path, media_type=document["mime"])


@app.get("/api/v1/admin/certificates")
def admin_certificates(_: None = Depends(require_admin)):
    return certificate_projection(limit=50)


@app.get("/api/v1/admin/catalog")
def catalog(_: None = Depends(require_admin)):
    return {
        "holders": fetch_all(
            """
            select id, concat(nombres, ' ', apellidos) as name,
                   concat(tipo_documento, ' ', documento_mascara) as document
            from titular
            order by created_at desc, nombres asc
            """
        ),
        "types": fetch_all(
            """
            select id, codigo as code, nombre as name, requiere_vigencia as requires_expiration
            from tipo_certificado
            where activo = true
            order by nombre asc
            """
        ),
    }


@app.get("/api/v1/admin/bulk/template")
def bulk_template(_: None = Depends(require_admin)):
    sample = io.StringIO()
    writer = csv.writer(sample, lineterminator="\n")
    writer.writerow(BULK_TEMPLATE_COLUMNS)
    writer.writerow(["Ana Maria", "Torres Ruiz", "DNI", "12345678", "ana@example.com", ""])
    writer.writerow(["Luis Alberto", "Ramos Perez", "DNI", "87654321", "luis@example.com", ""])
    return Response(
        content=sample.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="plantilla-carga-masiva-sicv.csv"'},
    )


@app.get("/api/v1/admin/bulk/template.xlsx")
def bulk_template_xlsx(_: None = Depends(require_admin)):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Participantes"
    sheet.append(BULK_TEMPLATE_COLUMNS)
    sheet.append(["Ana Maria", "Torres Ruiz", "DNI", "12345678", "ana@example.com", ""])
    sheet.append(["Luis Alberto", "Ramos Perez", "DNI", "87654321", "luis@example.com", ""])
    for cell in sheet[1]:
        cell.style = "Headline 3"
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla-carga-masiva-sicv.xlsx"'},
    )


@app.post("/api/v1/admin/bulk/preview")
def bulk_preview(
    file: UploadFile = File(...),
    type_id: str = Form(...),
    titulo: str = Form(...),
    programa: str | None = Form(default=None),
    fecha_emision: date = Form(...),
    vence_en: date | None = Form(default=None),
    origen: str = Form(default="NUEVO"),
    numero_prefijo: str = Form(default="SICV-2026-"),
    numero_inicio: int = Form(default=1),
    numero_modo: str = Form(default="AUTOMATICO"),
    _: None = Depends(require_admin),
):
    rows = read_bulk_file(file)
    preview = validate_bulk_rows(
        rows,
        type_id,
        titulo,
        fecha_emision,
        vence_en,
        origen,
        numero_prefijo.strip(),
        numero_inicio,
        numero_modo,
    )
    return {
        "total_rows": preview["total_rows"],
        "valid_rows": preview["valid_rows"],
        "error_rows": preview["error_rows"],
        "certificate_type": preview["certificate_type"]["nombre"],
        "rows": preview["rows"],
    }


@app.post("/api/v1/admin/bulk/import", status_code=status.HTTP_201_CREATED)
def bulk_import(
    file: UploadFile = File(...),
    type_id: str = Form(...),
    titulo: str = Form(...),
    programa: str | None = Form(default=None),
    fecha_emision: date = Form(...),
    vence_en: date | None = Form(default=None),
    origen: str = Form(default="NUEVO"),
    numero_prefijo: str = Form(default="SICV-2026-"),
    numero_inicio: int = Form(default=1),
    numero_modo: str = Form(default="AUTOMATICO"),
    allow_skip_invalid: bool = Form(default=False),
    actor: str = Form(default=SYSTEM_ACTOR),
    _: None = Depends(require_admin),
):
    rows = read_bulk_file(file)
    preview = validate_bulk_rows(
        rows,
        type_id,
        titulo,
        fecha_emision,
        vence_en,
        origen,
        numero_prefijo.strip(),
        numero_inicio,
        numero_modo,
    )
    return import_bulk_rows(
        preview,
        type_id,
        titulo,
        programa,
        fecha_emision,
        vence_en,
        origen,
        actor,
        allow_skip_invalid,
    )


@app.get("/api/v1/admin/bulk/history")
def bulk_history(_: None = Depends(require_admin)):
    return fetch_all(
        """
        select cm.id, cm.titulo, cm.programa, tc.nombre as certificate_type,
               cm.total_rows, cm.valid_rows, cm.error_rows,
               cm.created_count, cm.skipped_count, cm.actor, cm.status,
               to_char(cm.created_at, 'YYYY-MM-DD HH24:MI') as created_at
        from carga_masiva cm
        left join tipo_certificado tc on tc.id = cm.tipo_id
        where cm.institucion_id = %s
        order by cm.created_at desc
        limit 20
        """,
        (DEFAULT_INSTITUTION_ID,),
    )


@app.get("/api/v1/admin/bulk/{bulk_id}/details")
def bulk_details(bulk_id: str, _: None = Depends(require_admin)):
    bulk = fetch_one(
        """
        select cm.id, cm.titulo, cm.programa, tc.nombre as certificate_type,
               cm.total_rows, cm.valid_rows, cm.error_rows,
               cm.created_count, cm.skipped_count, cm.actor, cm.status,
               to_char(cm.created_at, 'YYYY-MM-DD HH24:MI') as created_at
        from carga_masiva cm
        left join tipo_certificado tc on tc.id = cm.tipo_id
        where cm.id = %s and cm.institucion_id = %s
        """,
        (bulk_id, DEFAULT_INSTITUTION_ID),
    )
    if bulk is None:
        raise HTTPException(status_code=404, detail="bulk_import_not_found")
    details = fetch_all(
        """
        select row_number, certificado_id, numero_institucional, titular,
               documento, status, errors, warnings, verification_url
        from carga_masiva_detalle
        where carga_id = %s
        order by row_number asc
        """,
        (bulk_id,),
    )
    return {"bulk": bulk, "details": details}


@app.post("/api/v1/admin/bulk/{bulk_id}/submit")
def bulk_submit_to_review(bulk_id: str, payload: StateAction, _: None = Depends(require_admin)):
    bulk = fetch_one(
        "select id from carga_masiva where id = %s and institucion_id = %s",
        (bulk_id, DEFAULT_INSTITUTION_ID),
    )
    if bulk is None:
        raise HTTPException(status_code=404, detail="bulk_import_not_found")
    return bulk_action_result(bulk_id, payload.actor)


@app.post("/api/v1/admin/holders", status_code=status.HTTP_201_CREATED)
def create_holder(payload: HolderCreate, _: None = Depends(require_admin)):
    normalized = normalize_document(payload.documento)
    try:
        holder = execute_one(
            """
            insert into titular (
              institucion_id, tipo_documento, documento_normalizado,
              documento_mascara, nombres, apellidos
            )
            values (%s, %s, %s, %s, %s, %s)
            returning id, concat(nombres, ' ', apellidos) as name,
                      concat(tipo_documento, ' ', documento_mascara) as document
            """,
            (
                DEFAULT_INSTITUTION_ID,
                payload.tipo_documento.upper(),
                normalized,
                mask_document(payload.documento),
                payload.nombres.strip(),
                payload.apellidos.strip(),
            ),
        )
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="holder_already_exists") from exc
    audit("CREAR_TITULAR", holder["name"])
    return holder


@app.post("/api/v1/admin/certificates", status_code=status.HTTP_201_CREATED)
def create_certificate(payload: CertificateCreate, _: None = Depends(require_admin)):
    if payload.vence_en and payload.vence_en < payload.fecha_emision:
        raise HTTPException(status_code=422, detail="expires_before_issue_date")
    if payload.origen not in {"NUEVO", "HISTORICO"}:
        raise HTTPException(status_code=422, detail="invalid_origin")

    token = generate_public_token()
    try:
        certificate = execute_one(
            """
            insert into certificado (
              institucion_id, titular_id, tipo_id, numero_institucional,
              token_publico, estado, origen, titulo, programa, fecha_emision, vence_en
            )
            values (%s, %s, %s, %s, %s, 'BORRADOR', %s, %s, %s, %s, %s)
            returning id
            """,
            (
                DEFAULT_INSTITUTION_ID,
                payload.holder_id,
                payload.type_id,
                payload.numero_institucional.strip(),
                token,
                payload.origen,
                payload.titulo.strip(),
                payload.programa.strip() if payload.programa else None,
                payload.fecha_emision,
                payload.vence_en,
            ),
        )
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="certificate_number_exists") from exc
    except psycopg.errors.ForeignKeyViolation as exc:
        raise HTTPException(status_code=422, detail="invalid_holder_or_type") from exc

    generate_qr_svg(str(certificate["id"]), token)
    audit("CREAR_CERTIFICADO", payload.numero_institucional)
    audit("GENERAR_QR", payload.numero_institucional)
    return certificate_projection("where c.id = %s", (certificate["id"],))[0]


@app.post("/api/v1/admin/certificates/{certificate_id}/documents", status_code=status.HTTP_201_CREATED)
def upload_certificate_document(certificate_id: str, file: UploadFile = File(...), _: None = Depends(require_admin)):
    certificate = get_certificate_or_404(certificate_id)
    if file.content_type not in ALLOWED_DOCUMENT_MIME:
        raise HTTPException(status_code=415, detail="unsupported_document_type")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    extension = {
        "application/pdf": ".pdf",
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
    }[file.content_type]
    object_key = f"{certificate_id}{extension}"
    destination = UPLOAD_DIR / object_key
    digest = hashlib.sha256()
    size = 0

    with destination.open("wb") as output:
        while chunk := file.file.read(1024 * 1024):
            size += len(chunk)
            if size > MAX_DOCUMENT_BYTES:
                output.close()
                destination.unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="document_too_large")
            digest.update(chunk)
            output.write(chunk)

    with db_cursor() as cur:
        cur.execute(
            "update documento_version set es_definitivo = false where certificado_id = %s",
            (certificate_id,),
        )
        cur.execute(
            """
            select coalesce(max(version), 0) + 1 as next_version
            from documento_version
            where certificado_id = %s
            """,
            (certificate_id,),
        )
        version = cur.fetchone()["next_version"]
        cur.execute(
            """
            insert into documento_version (
              certificado_id, version, object_key, sha256, mime, size_bytes, es_definitivo
            )
            values (%s, %s, %s, %s, %s, %s, true)
            returning id, version, object_key, sha256, mime, size_bytes, es_definitivo
            """,
            (certificate_id, version, object_key, digest.hexdigest(), file.content_type, size),
        )
        document = cur.fetchone()

    audit("SUBIR_DOCUMENTO", certificate["numero_institucional"])
    return document


@app.get("/api/v1/admin/certificates/{certificate_id}/document")
def admin_certificate_document(certificate_id: str, _: None = Depends(require_admin)):
    get_certificate_or_404(certificate_id)
    document = fetch_one(
        """
        select object_key, mime
        from documento_version
        where certificado_id = %s and es_definitivo = true
        order by version desc
        limit 1
        """,
        (certificate_id,),
    )
    if document is None:
        raise HTTPException(status_code=404, detail="document_not_found")
    path = UPLOAD_DIR / document["object_key"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="document_file_missing")
    return FileResponse(path, media_type=document["mime"])


@app.get("/api/v1/admin/certificates/{certificate_id}/qr")
def admin_certificate_qr(certificate_id: str, _: None = Depends(require_admin)):
    certificate = get_certificate_or_404(certificate_id)
    qr = ensure_qr_for_certificate(certificate)
    path = UPLOAD_DIR / qr["object_key"]
    if not path.exists():
        qr = generate_qr_svg(str(certificate["id"]), certificate["token_publico"])
        path = UPLOAD_DIR / qr["object_key"]
    return FileResponse(
        path,
        media_type=qr["mime"],
        filename=f"qr-{certificate['numero_institucional']}.svg",
    )


@app.get("/api/v1/admin/certificates/{certificate_id}/qr-metadata")
def admin_certificate_qr_metadata(certificate_id: str, _: None = Depends(require_admin)):
    certificate = get_certificate_or_404(certificate_id)
    qr = ensure_qr_for_certificate(certificate)
    return {
        "token": certificate["token_publico"],
        "verification_url": qr["verification_url"],
        "sha256": qr["sha256"],
        "mime": qr["mime"],
    }


@app.get("/api/v1/admin/review-queue")
def review_queue(_: None = Depends(require_admin)):
    return fetch_all(
        """
        select c.id,
               c.numero_institucional as code,
               concat(t.nombres, ' ', t.apellidos) as holder,
               c.titulo as title,
               c.origen as origin,
               c.estado as status,
               case
                 when c.updated_at > now() - interval '1 minute' then 'Hace instantes'
                 when c.updated_at > now() - interval '1 hour' then 'Hace ' || floor(extract(epoch from now() - c.updated_at) / 60)::int || ' min'
                 when c.updated_at > now() - interval '1 day' then 'Hace ' || floor(extract(epoch from now() - c.updated_at) / 3600)::int || ' h'
                 else to_char(c.updated_at, 'YYYY-MM-DD')
               end as updated,
               case
                 when pd.id is null then 'Sin politica activa'
                 when c.estado = 'OBSERVADO' then 'Requiere correccion'
                 when c.estado = 'APROBADO' then 'Listo para publicar'
                 else 'Politica activa'
               end as risk
        from certificado c
        join titular t on t.id = c.titular_id
        left join politica_divulgacion pd on pd.institucion_id = c.institucion_id
          and pd.activa = true
          and (pd.tipo_id = c.tipo_id or pd.tipo_id is null)
        where c.estado in ('BORRADOR', 'EN_REVISION', 'OBSERVADO', 'APROBADO')
        order by c.updated_at desc
        """
    )


@app.post("/api/v1/admin/certificates/{certificate_id}/submit")
def submit_certificate(certificate_id: str, payload: StateAction, _: None = Depends(require_admin)):
    return transition_certificate(
        certificate_id,
        {"BORRADOR", "OBSERVADO"},
        "EN_REVISION",
        "ENVIAR_REVISION",
        payload.actor,
        payload.reason,
    )


@app.post("/api/v1/admin/certificates/{certificate_id}/observe")
def observe_certificate(certificate_id: str, payload: StateAction, _: None = Depends(require_admin)):
    if not payload.reason:
        raise HTTPException(status_code=422, detail="reason_required")
    return transition_certificate(
        certificate_id,
        {"EN_REVISION"},
        "OBSERVADO",
        "OBSERVAR",
        payload.actor,
        payload.reason,
    )


@app.post("/api/v1/admin/certificates/{certificate_id}/approve")
def approve_certificate(certificate_id: str, payload: StateAction, _: None = Depends(require_admin)):
    return transition_certificate(
        certificate_id,
        {"EN_REVISION"},
        "APROBADO",
        "APROBAR",
        payload.actor,
        payload.reason,
    )


@app.post("/api/v1/admin/certificates/{certificate_id}/publish")
def publish_certificate(certificate_id: str, payload: StateAction, _: None = Depends(require_admin)):
    certificate = get_certificate_or_404(certificate_id)
    if certificate["estado"] != "APROBADO":
        raise HTTPException(status_code=409, detail="invalid_state_transition")

    policy = fetch_one(
        """
        select id from politica_divulgacion
        where institucion_id = %s and activa = true
        limit 1
        """,
        (certificate["institucion_id"],),
    )
    if policy is None:
        raise HTTPException(status_code=409, detail="active_disclosure_policy_required")

    qr = ensure_qr_for_certificate(certificate)
    token = certificate["token_publico"]
    execute_one(
        """
        update certificado
        set estado = 'PUBLICADO', token_publico = %s, revisor = %s, updated_at = now()
        where id = %s
        returning id
        """,
        (token, payload.actor, certificate_id),
    )
    audit("PUBLICAR", certificate["numero_institucional"], payload.actor)
    audit("CONFIRMAR_QR", qr["verification_url"], payload.actor)
    return certificate_projection("where c.id = %s", (certificate_id,))[0]


@app.post("/api/v1/admin/certificates/{certificate_id}/revoke")
def revoke_certificate(certificate_id: str, payload: StateAction, _: None = Depends(require_admin)):
    if not payload.reason:
        raise HTTPException(status_code=422, detail="reason_required")
    return transition_certificate(
        certificate_id,
        {"PUBLICADO"},
        "REVOCADO",
        "REVOCAR",
        payload.actor,
        payload.reason,
    )


@app.get("/api/v1/admin/audit-events")
def audit_events(_: None = Depends(require_admin)):
    return fetch_all(
        """
        select to_char(creado_en, 'HH24:MI') as time, accion as action,
               objetivo as target, actor
        from evento_auditoria
        order by creado_en desc
        limit 20
        """
    )


@app.get("/api/v1/admin/disclosure-policy")
def disclosure_policy(_: None = Depends(require_admin)):
    policy = fetch_one(
        """
        select version, descripcion, mostrar_codigo, mostrar_nombre, mostrar_documento,
               mostrar_fechas, mostrar_programa, permitir_documento_completo
        from politica_divulgacion
        where activa = true
        order by tipo_id nulls last
        limit 1
        """
    )
    if policy is None:
        raise HTTPException(status_code=404, detail="policy_not_found")
    return policy


@app.post("/api/v1/admin/disclosure-policy", status_code=status.HTTP_201_CREATED)
def create_disclosure_policy(payload: PolicyCreate, _: None = Depends(require_admin)):
    current = fetch_one(
        """
        select coalesce(max(version), 0) + 1 as next_version
        from politica_divulgacion
        where institucion_id = %s and tipo_id is null
        """,
        (DEFAULT_INSTITUTION_ID,),
    )
    with db_cursor() as cur:
        cur.execute(
            """
            update politica_divulgacion
            set activa = false
            where institucion_id = %s and tipo_id is null and activa = true
            """,
            (DEFAULT_INSTITUTION_ID,),
        )
        cur.execute(
            """
            insert into politica_divulgacion (
              institucion_id, version, descripcion, activa, mostrar_codigo,
              mostrar_nombre, mostrar_documento, mostrar_fechas, mostrar_programa,
              permitir_documento_completo, approved_at
            )
            values (%s, %s, %s, true, %s, %s, %s, %s, %s, %s, now())
            returning version, descripcion, mostrar_codigo, mostrar_nombre,
                      mostrar_documento, mostrar_fechas, mostrar_programa,
                      permitir_documento_completo
            """,
            (
                DEFAULT_INSTITUTION_ID,
                current["next_version"],
                payload.descripcion,
                payload.mostrar_codigo,
                payload.mostrar_nombre,
                payload.mostrar_documento,
                payload.mostrar_fechas,
                payload.mostrar_programa,
                payload.permitir_documento_completo,
            ),
        )
        policy = cur.fetchone()
    audit("ACTIVAR_POLITICA", f"Divulgacion v{policy['version']}", payload.actor)
    return policy


@app.get("/api/v1/admin/summary")
def summary(_: None = Depends(require_admin)):
    row = fetch_one(
        """
        select
          count(*) filter (where estado = 'PUBLICADO') as published,
          count(*) filter (where estado = 'REVOCADO') as revoked,
          count(*) filter (where estado in ('BORRADOR', 'EN_REVISION', 'OBSERVADO', 'APROBADO')) as in_review,
          (select count(*) from politica_divulgacion where activa = true) as active_policies
        from certificado
        """
    )
    return row


@app.get("/api/v1/admin/settings")
def get_admin_settings(_: None = Depends(require_admin)):
    institution = fetch_one(
        "select id, razon_social, dominio_verificacion, activo, created_at from institucion where id = %s",
        (DEFAULT_INSTITUTION_ID,),
    )
    if institution is None:
        execute_one(
            """
            insert into institucion (id, razon_social, dominio_verificacion)
            values (%s, 'Institucion Educativa Demo', 'localhost')
            on conflict (id) do nothing
            """,
            (DEFAULT_INSTITUTION_ID,),
        )
        institution = fetch_one(
            "select id, razon_social, dominio_verificacion, activo, created_at from institucion where id = %s",
            (DEFAULT_INSTITUTION_ID,),
        )

    effective_base_url = get_verification_base_url()
    sample_token = "SICV-DEMO-001"

    certificates_count_row = fetch_one(
        "select count(*) as total, count(*) filter (where estado = 'PUBLICADO') as published from certificado"
    )

    return {
        "institution": institution,
        "effective_verification_base_url": effective_base_url,
        "env_verification_base_url": PUBLIC_VERIFICATION_BASE_URL,
        "sample_qr_url": f"{effective_base_url}/{sample_token}",
        "certificates_count": certificates_count_row["total"] if certificates_count_row else 0,
        "published_count": certificates_count_row["published"] if certificates_count_row else 0,
        "is_production_mode": not ("localhost" in effective_base_url or "127.0.0.1" in effective_base_url),
    }


@app.put("/api/v1/admin/settings")
def update_admin_settings(payload: InstitutionSettingsUpdate, _: None = Depends(require_admin)):
    cleaned_razon = payload.razon_social.strip()
    cleaned_domain = payload.dominio_verificacion.strip()

    with db_cursor() as cur:
        cur.execute(
            """
            update institucion
            set razon_social = %s,
                dominio_verificacion = %s
            where id = %s
            returning id, razon_social, dominio_verificacion, activo
            """,
            (cleaned_razon, cleaned_domain, DEFAULT_INSTITUTION_ID),
        )
        updated = cur.fetchone()

    audit("ACTUALIZAR_CONFIGURACION", f"Institucion: {cleaned_razon} | Dominio: {cleaned_domain}")
    effective_base_url = get_verification_base_url()

    return {
        "institution": updated,
        "effective_verification_base_url": effective_base_url,
        "sample_qr_url": f"{effective_base_url}/SICV-DEMO-001",
        "is_production_mode": not ("localhost" in effective_base_url or "127.0.0.1" in effective_base_url),
        "message": "Configuracion institucional actualizada exitosamente.",
    }


@app.post("/api/v1/admin/settings/regenerate-qrs")
def regenerate_all_qrs(_: None = Depends(require_admin)):
    certificates = fetch_all(
        "select id, token_publico, numero_institucional from certificado where token_publico is not null"
    )
    count = 0
    with db_cursor() as cur:
        for cert in certificates:
            generate_qr_svg_in_cursor(cur, str(cert["id"]), cert["token_publico"])
            count += 1

    audit("REGENERAR_QRS", f"Se regeneraron {count} codigos QR con el dominio actual.")
    return {
        "status": "success",
        "regenerated_count": count,
        "effective_verification_base_url": get_verification_base_url(),
        "message": f"Se han regenerado exitosamente {count} codigos QR.",
    }


@app.post("/api/v1/admin/system/reset-workspace")
def reset_workspace(_: None = Depends(require_admin)):
    with db_cursor() as cur:
        cur.execute("delete from documento_version")
        cur.execute("delete from certificado_qr")
        cur.execute("delete from carga_masiva_detalle")
        cur.execute("delete from carga_masiva")
        cur.execute("delete from certificado")
        cur.execute("delete from titular")
        cur.execute("delete from cola_revision")
        cur.execute("delete from evento_auditoria")
        cur.execute("delete from politica_divulgacion")
        cur.execute(
            """
            insert into politica_divulgacion (
              id, institucion_id, tipo_id, version, descripcion, activa,
              mostrar_codigo, mostrar_nombre, mostrar_documento, mostrar_fechas,
              mostrar_programa, permitir_documento_completo, approved_at
            )
            values (
              '44444444-4444-4444-4444-444444444441',
              %s,
              null,
              1,
              'Codigo, titulo, institucion, fecha, programa y estado. Documento completo oculto por defecto.',
              true, true, true, false, true, true, false, now()
            )
            """,
            (DEFAULT_INSTITUTION_ID,),
        )
    clean_uploads()
    return {"status": "clean", "message": "workspace_reset_complete"}

